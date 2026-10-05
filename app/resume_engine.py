from __future__ import annotations

import re
from pathlib import Path
import hashlib

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


STANDARD_HEADINGS = {
    "PROFESSIONAL SUMMARY",
    "SUMMARY",
    "SKILLS",
    "TECHNICAL SKILLS",
    "WORK EXPERIENCE",
    "EXPERIENCE",
    "PROJECTS",
    "EDUCATION",
    "CERTIFICATIONS",
    "CERTIFICATIONS AND TRAINING",
}

BULLET_PREFIXES = ("-", "•", "*", "▪")


def _clean_filename(value: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip())
    return value.strip("._")[:100] or "resume"


def _is_heading(line: str) -> bool:
    normalized = re.sub(r"[:\s]+$", "", line.strip()).upper()
    return normalized in STANDARD_HEADINGS


def _add_run(paragraph, text: str, bold: bool = False, italic: bool = False) -> None:
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.name = "Arial"
    run.font.size = Pt(10.5)


class ATSResumeEngine:
    """Render a plain, single-column ATS-friendly DOCX resume."""

    def render_docx(
        self,
        resume_text: str,
        company: str,
        role: str,
        output_dir: str = "./data/applications",
    ) -> str:
        out_dir = Path(output_dir)
        out_dir.mkdir(parents=True, exist_ok=True)

        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)

        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(10.5)
        normal.paragraph_format.space_after = Pt(2)
        normal.paragraph_format.line_spacing = 1.0

        lines = [line.rstrip() for line in resume_text.splitlines()]
        first_nonempty = next((line.strip() for line in lines if line.strip()), "")

        start_index = 0
        if first_nonempty and not _is_heading(first_nonempty):
            p = document.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(first_nonempty)
            run.bold = True
            run.font.name = "Arial"
            run.font.size = Pt(16)
            start_index = lines.index(next(line for line in lines if line.strip())) + 1

            contact_lines = []
            while start_index < len(lines) and lines[start_index].strip() and not _is_heading(lines[start_index]):
                contact_lines.append(lines[start_index].strip())
                start_index += 1
            if contact_lines:
                p = document.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                _add_run(p, " | ".join(contact_lines))

        current_heading = None
        for line in lines[start_index:]:
            stripped = line.strip()
            if not stripped:
                continue

            if _is_heading(stripped):
                current_heading = stripped.rstrip(":").upper()
                p = document.add_paragraph()
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(2)
                run = p.add_run(current_heading)
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(11.5)
                continue

            p = document.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18) if stripped.startswith(BULLET_PREFIXES) else Inches(0)
            p.paragraph_format.first_line_indent = Inches(-0.12) if stripped.startswith(BULLET_PREFIXES) else Inches(0)
            text = stripped
            if stripped.startswith(BULLET_PREFIXES):
                text = stripped[1:].strip()
                _add_run(p, "• ", bold=False)
            _add_run(p, text)

        job_suffix = hashlib.sha1(f"{company}|{role}".encode("utf-8")).hexdigest()[:8]
        output = out_dir / f"{_clean_filename(company)}_{_clean_filename(role)}_{job_suffix}_resume.docx"
        document.core_properties.title = f"{role} Resume"
        document.core_properties.subject = f"ATS resume for {company}"
        document.core_properties.author = ""
        document.save(output)
        return str(output)

    @staticmethod
    def extract_text(docx_path: str) -> str:
        doc = Document(docx_path)
        return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
